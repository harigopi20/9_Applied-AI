#!/usr/bin/env python3
"""Build the single sample data file app/data/baumarkt-buddy-data.xlsx.

DR-01..DR-03: ONE Excel file holds products, prices, stores, stock and
aisle/shelf per store. This script converts the team's CSV files
(data/product_catalogue.csv, data/store_map.csv) into that workbook and
adds the columns the rule engine needs (variant_group, pack_qty, specs).

All values are SAMPLE DATA (DR-02, NFR-05). Store S01 uses the original
store_map.csv; stores S02/S03 are deliberately different so that RB-15
(only the selected store's data) and RB-06 (missing aisle never invented)
can be tested.

Usage:  python3 tools/build_sample_data.py
"""
import csv
import re
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
SRC_PRODUCTS = ROOT / "data" / "product_catalogue.csv"
SRC_STOCK = ROOT / "data" / "store_map.csv"
OUT = ROOT / "app" / "data" / "baumarkt-buddy-data.xlsx"
PLZ_TABLE = ROOT / "app" / "data" / "plz-de.csv"  # built by tools/build_plz_table.py

PRODUCT_COLUMNS = [
    "product_nr", "product_name", "variant_group", "category", "size",
    "pack_qty", "pack_unit", "specs", "description", "material", "colour",
    "price_eur", "max_load_kg", "suitable_wall_types", "compatible_with",
    "use_cases", "requires_professional", "safety_note", "data_status",
    "data_note",
]

# Structured specs used by the compatibility checks (FR-09, RB-03).
# Derived from the "size" text of the CSV; key=value pairs separated by ';'.
SPECS = {
    "P035": "diameter_mm=6", "P036": "diameter_mm=8", "P037": "diameter_mm=10",
    "P043": "diameter_mm=6", "P044": "diameter_mm=8", "P045": "diameter_mm=10",
    "P046": "thread=M5", "P047": "thread=M8;length_mm=80",
    "P048": "diameter_mm=4;length_mm=40", "P049": "diameter_mm=5;length_mm=50",
    "P050": "diameter_mm=6;length_mm=60",
    "P058": "arm_mm=150", "P059": "arm_mm=200", "P060": "arm_mm=250",
    "P063": "length_mm=800;depth_mm=200;thickness_mm=18",
    "P064": "length_mm=800;depth_mm=250;thickness_mm=18",
}

# Additions to the team CSV, documented per product in data_note.
OVERRIDES = {
    "P051": {
        "suitable_wall_types": "brick;plasterboard",
        "data_note": "suitable_wall_types added for the sample file - team to verify (SR-10)",
    },
}


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def pack_info(row):
    """pack_qty = number of single pieces in one sales unit.

    Only filled where the size text states it clearly ("50 pcs").
    Sets such as '3 pcs 30/50/70 mm' (brush set) are ONE sales unit.
    """
    m = re.search(r"(\d+)\s*pcs\b", row["size"])
    if m and row["category"] in ("Fixings", "Shelving", "Safety"):
        return int(m.group(1)), "pcs"
    return 1, "item"


def num(value):
    if value in ("", None):
        return None
    f = float(value)
    return int(f) if f.is_integer() else f


def main():
    with SRC_PRODUCTS.open(encoding="utf-8") as fh:
        products = list(csv.DictReader(fh))
    with SRC_STOCK.open(encoding="utf-8") as fh:
        base_stock = {r["product_nr"]: r for r in csv.DictReader(fh)}

    wb = openpyxl.Workbook()
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="FFE699")

    # ---- Info sheet -------------------------------------------------------
    info = wb.active
    info.title = "Info"
    for line in [
        ["Baumarkt-Buddy sample data (R1 MVP)"],
        ["SAMPLE DATA - prices, stock, stores, aisles and shelves are invented for the prototype (DR-02, NFR-05)."],
        ["Built from data/product_catalogue.csv and data/store_map.csv by tools/build_sample_data.py"],
        ["Build date", date.today().isoformat()],
        ["Sheets", "Products, Stores, Stock"],
        ["Stock: one row per store and product. Empty aisle/shelf = not known; the app shows 'not available' and never guesses (RB-06)."],
    ]:
        info.append(line)
    info["A1"].font = Font(bold=True, size=14)

    # ---- Products ---------------------------------------------------------
    ws = wb.create_sheet("Products")
    ws.append(PRODUCT_COLUMNS)
    for p in products:
        qty, unit = pack_info(p)
        row = {
            "product_nr": p["product_nr"],
            "product_name": p["product_name"],
            "variant_group": slug(p["product_name"]),
            "category": p["category"],
            "size": p["size"],
            "pack_qty": qty,
            "pack_unit": unit,
            "specs": SPECS.get(p["product_nr"], ""),
            "description": p["description"],
            "material": p["material"],
            "colour": p["colour"],
            "price_eur": num(p["price_eur"]),
            "max_load_kg": num(p["max_load_kg"]),
            "suitable_wall_types": p["suitable_wall_types"],
            "compatible_with": p["compatible_with"],
            "use_cases": p["use_cases"],
            "requires_professional": p["requires_professional"],
            "safety_note": p["safety_note"],
            "data_status": "sample",
            "data_note": "",
        }
        row.update(OVERRIDES.get(p["product_nr"], {}))
        ws.append([row[c] for c in PRODUCT_COLUMNS])

    # ---- Stores -----------------------------------------------------------
    # Coordinates are NOT invented: they are the centre point of the store's
    # PLZ from the PLZ reference table app/data/plz-de.csv (DR-08, OpenGeoDB).
    with PLZ_TABLE.open(encoding="utf-8") as fh:
        plz_coords = {r["plz"]: (float(r["lat"]), float(r["lon"])) for r in csv.DictReader(fh)}
    stores = wb.create_sheet("Stores")
    stores.append(["store_id", "store_name", "plz", "city", "lat", "lon", "data_status"])
    for sid, name, plz, city in [
        ("S01", "Sample Store Berlin-Mitte", "10117", "Berlin"),
        ("S02", "Sample Store München-Zentrum", "80331", "München"),
        ("S03", "Sample Store Hamburg-Altstadt", "20095", "Hamburg"),
    ]:
        lat, lon = plz_coords[plz]
        stores.append([sid, name, plz, city, lat, lon, "sample"])

    # ---- Stock (per store) ------------------------------------------------
    stock = wb.create_sheet("Stock")
    stock.append(["store_id", "product_nr", "stock_qty", "aisle_nr", "aisle_name", "shelf_nr", "data_status"])
    # S01: exactly the team's store_map.csv
    for p in products:
        b = base_stock[p["product_nr"]]
        stock.append(["S01", p["product_nr"], num(b["qty_on_shelf"]), num(b["aisle_nr"]),
                      b["aisle_name"], num(b["shelf_nr"]), "sample"])
    # S02: different layout (aisle + 20), different stock; P036 and P049
    # have NO aisle/shelf data (RB-06), P059 is out of stock.
    for i, p in enumerate(products):
        b = base_stock[p["product_nr"]]
        nr = p["product_nr"]
        qty = max(0, int(b["qty_on_shelf"]) - 3 - (i % 5))
        aisle, name, shelf = int(b["aisle_nr"]) + 20, b["aisle_name"], int(b["shelf_nr"])
        if nr in ("P036", "P049"):
            aisle, name, shelf = None, "", None
        if nr == "P059":
            qty = 0
        stock.append(["S02", nr, qty, aisle, name, shelf, "sample"])
    # S03: another layout (aisle reversed); P060 and P061 are not stocked
    # at all (no row -> stock and location unknown).
    for i, p in enumerate(products):
        b = base_stock[p["product_nr"]]
        if p["product_nr"] in ("P060", "P061"):
            continue
        stock.append(["S03", p["product_nr"], int(b["qty_on_shelf"]) + (i % 7),
                      40 - int(b["aisle_nr"]), b["aisle_name"], int(b["shelf_nr"]) + 1, "sample"])

    for sheet in (ws, stores, stock):
        for cell in sheet[1]:
            cell.font = head_font
            cell.fill = head_fill
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col[:80])
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"written {OUT.relative_to(ROOT)}: {len(products)} products, 3 stores")


if __name__ == "__main__":
    main()
